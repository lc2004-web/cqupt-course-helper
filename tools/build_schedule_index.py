#!/usr/bin/env python3
"""Build the extension's compact course/class catalog from CQUPT schedule workbooks."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_HEADERS = (
    "开课院系",
    "课程编号",
    "课程名称",
    "班级名称",
    "任课教师",
    "上课信息描述",
    "上课地点",
)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def teacher_names(value: str) -> list[str]:
    names: list[str] = []
    for part in re.split(r"[；;、]+", value):
        name = clean(part.split("|", 1)[-1])
        name = re.sub(r"[（(](?:主讲|参讲)[）)]$", "", name).strip()
        if name and name not in names:
            names.append(name)
    return names


def natural_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", value)]


def read_records(path: Path) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        headers: list[str] | None = None
        for row in rows:
            candidate = [clean(value) for value in row]
            if "课程编号" in candidate and "班级名称" in candidate:
                headers = candidate
                break
        if not headers:
            continue
        indexes = {header: headers.index(header) for header in REQUIRED_HEADERS}
        for row in rows:
            code = clean(row[indexes["课程编号"]]).upper()
            if not re.fullmatch(r"[A-Z0-9][A-Z0-9_-]{2,30}", code):
                continue
            teacher_text = clean(row[indexes["任课教师"]])
            records.append(
                {
                    "code": code,
                    "courseName": clean(row[indexes["课程名称"]]),
                    "className": clean(row[indexes["班级名称"]]),
                    "teachers": teacher_names(teacher_text),
                    "teacherText": teacher_text,
                    "schedule": clean(row[indexes["上课信息描述"]]),
                    "location": clean(row[indexes["上课地点"]]),
                    "department": clean(row[indexes["开课院系"]]),
                }
            )
    return records


def build_catalog(inputs: list[Path]) -> dict[str, object]:
    unique: dict[tuple[str, ...], dict[str, object]] = {}
    for path in inputs:
        for record in read_records(path):
            key = (
                str(record["code"]),
                str(record["className"]),
                str(record["teacherText"]),
                str(record["schedule"]),
                str(record["location"]),
            )
            unique[key] = record

    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    names: dict[str, Counter[str]] = defaultdict(Counter)
    for record in unique.values():
        code = str(record["code"])
        grouped[code].append(record)
        names[code][str(record["courseName"])] += 1

    courses: dict[str, object] = {}
    for code in sorted(grouped):
        classes = sorted(grouped[code], key=lambda item: natural_key(str(item["className"])))
        normalized_classes = []
        for item in classes:
            normalized_classes.append(
                {
                    "id": f'{code}:{item["className"]}',
                    "className": item["className"],
                    "teachers": item["teachers"],
                    "teacherText": item["teacherText"],
                    "schedule": item["schedule"],
                    "location": item["location"],
                    "department": item["department"],
                }
            )
        courses[code] = {
            "name": names[code].most_common(1)[0][0],
            "classes": normalized_classes,
        }

    return {
        "catalogVersion": "2026-2027-1",
        "schoolYear": "2026-2027",
        "term": "第一学期",
        "sourceFiles": [path.name for path in inputs],
        "courseCount": len(courses),
        "classCount": len(unique),
        "courses": courses,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    catalog = build_catalog(args.inputs)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(catalog, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(
        f'Wrote {catalog["courseCount"]} courses and '
        f'{catalog["classCount"]} classes to {args.output}'
    )


if __name__ == "__main__":
    main()
